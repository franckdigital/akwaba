import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.core import audit
from apps.core.permissions import allowed

from . import services


class DashboardView(APIView):
    def get(self, request):
        return Response(services.dashboard(request.user, request.query_params))


class KpiView(APIView):
    def get(self, request):
        if not allowed(request.user, "reports", "view"):
            raise PermissionDenied()
        return Response(services.kpis(request.user, request.query_params))


class ProfitabilityView(APIView):
    def get(self, request):
        if not allowed(request.user, "reports", "view") or request.user.role in ("learner", "instructor"):
            raise PermissionDenied()
        return Response(services.profitability(request.user, request.query_params))


class MonthlyFinanceView(APIView):
    """GET /api/reports/monthly-finance/?date_from=&date_to=  ->  {headers, rows} (dernière ligne = TOTAL)."""

    def get(self, request):
        if not allowed(request.user, "reports", "view"):
            raise PermissionDenied("Accès non autorisé.")
        headers, rows = services.monthly_finance(request.user, request.query_params)
        return Response({"headers": headers, "rows": rows})


class ReportListView(APIView):
    def get(self, request):
        return Response([{"key": k, "title": t} for k, (t, _, res) in services.REPORTS.items()
                         if allowed(request.user, res, "view") and allowed(request.user, "reports", "view")
                         or (request.user.role == "learner" and k in ("payments", "quiz_results", "exams", "certificates", "progress"))])


def _fmt(v):
    return "" if v is None else v


def _pdf(title, headers, rows, subtitle=""):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas
    buf = BytesIO()
    w, h = landscape(A4)
    c = canvas.Canvas(buf, pagesize=(w, h))
    y = h - 50
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, title)
    y -= 16
    if subtitle:
        c.setFont("Helvetica", 9)
        c.drawString(40, y, subtitle)
    y -= 20
    col = (w - 80) / max(len(headers), 1)

    def head():
        c.setFont("Helvetica-Bold", 8)
        for i, hd in enumerate(headers):
            c.drawString(40 + i * col, y, str(hd)[:22])
    head()
    c.setFont("Helvetica", 8)
    for row in rows:
        y -= 14
        if y < 40:
            c.showPage()
            y = h - 50
            head()
            c.setFont("Helvetica", 8)
            y -= 14
        bold = bool(row) and str(row[0]).upper() == "TOTAL"
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 8)
        for i, v in enumerate(row):
            c.drawString(40 + i * col, y, str(_fmt(v))[:24])
    c.showPage()
    c.save()
    return buf.getvalue()


class ExportView(APIView):
    """GET /api/reports/export/?report=payments&format=xlsx|csv|pdf[&date_from=&date_to=&school=]"""

    def get(self, request):
        name = request.query_params.get("report")
        fmt = request.query_params.get("format", "csv")
        if name not in services.REPORTS:
            raise ValidationError({"report": "Rapport inconnu. Disponibles : " + ", ".join(services.REPORTS)})
        title, fn, resource = services.REPORTS[name]
        u = request.user
        if not allowed(u, resource, "export") and not (u.role == "learner" and name in ("payments", "quiz_results", "exams", "certificates", "progress")):
            raise PermissionDenied("Export non autorisé.")
        headers, rows = fn(u, request.query_params)
        audit.log(request, "export", model="reports", new={"report": name, "format": fmt, "rows": len(rows)})
        if fmt == "csv":
            out = StringIO()
            wr = csv.writer(out, delimiter=";")
            wr.writerow(headers)
            wr.writerows([[_fmt(v) for v in r] for r in rows])
            resp = HttpResponse("﻿" + out.getvalue(), content_type="text/csv; charset=utf-8")
        elif fmt == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = title[:30]
            ws.append(headers)
            for r in rows:
                ws.append([_fmt(v) if not hasattr(v, "isoformat") or isinstance(v, str) else v.isoformat() for v in r])
            from openpyxl.styles import Font, PatternFill
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="E02027")
            for i, hd in enumerate(headers, start=1):
                width = max([len(str(hd))] + [len(str(_fmt(r[i - 1]))) for r in rows[:200] if len(r) >= i]) + 2
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 48)
            if rows and str(rows[-1][0]).upper() == "TOTAL":
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
            ws.freeze_panes = "A2"
            resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            wb.save(resp)
        elif fmt == "pdf":
            sub = "Période : %s → %s" % (request.query_params.get("date_from") or "début", request.query_params.get("date_to") or "aujourd'hui")
            resp = HttpResponse(_pdf(title, headers, rows, sub), content_type="application/pdf")
        else:
            raise ValidationError({"format": "csv, xlsx ou pdf."})
        resp["Content-Disposition"] = f'attachment; filename="{name}.{fmt}"'
        return resp
