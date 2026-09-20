from datetime import date, timedelta

from django.db import transaction
from django.http import HttpResponse
from rest_framework import serializers
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response

from apps.core import audit
from apps.core.permissions import ORG_ADMIN, ORG_ROLES
from apps.core.viewsets import SCHOOL, BaseModelSerializer, ScopedModelViewSet, make_serializer
from apps.learners.models import Learner

from .models import Cohort, Contract, Group, Organization, QuoteRequest

OrganizationSerializer = make_serializer(Organization, extra_fields={
    "learners_count": serializers.SerializerMethodField(),
    "get_learners_count": lambda self, o: o.learners.count()})
GroupSerializer = make_serializer(Group, extra_fields={
    "learners_count": serializers.SerializerMethodField(),
    "get_learners_count": lambda self, o: o.learners.count()})


class OrganizationViewSet(ScopedModelViewSet):
    resource = "organizations"
    queryset = Organization.objects.all()
    serializer_class = OrganizationSerializer
    org_lookup = "id"
    filterset_fields = ["org_type", "is_active", "school"]
    search_fields = ["name", "contact_name", "email"]


class ContractSerializer(BaseModelSerializer):
    net_price = serializers.DecimalField(max_digits=12, decimal_places=0, read_only=True)
    learner_contribution = serializers.DecimalField(max_digits=12, decimal_places=0, read_only=True)
    organization_total = serializers.DecimalField(max_digits=14, decimal_places=0, read_only=True)
    organization_name = serializers.CharField(source="organization.name", read_only=True)
    training_name = serializers.CharField(source="training.name", read_only=True)

    class Meta:
        model = Contract
        fields = "__all__"

    def validate(self, attrs):
        org = attrs.get("organization") or getattr(self.instance, "organization", None)
        training = attrs.get("training") or getattr(self.instance, "training", None)
        if org and training and org.school_id != training.school_id:
            raise serializers.ValidationError("L'organisation et la formation doivent relever de la même auto-école.")
        price = attrs.get("unit_price", getattr(self.instance, "unit_price", 0))
        contrib = attrs.get("organization_contribution", getattr(self.instance, "organization_contribution", 0))
        discount = attrs.get("discount_percent", getattr(self.instance, "discount_percent", 0))
        net = price * (100 - discount) / 100
        if contrib > net:
            raise serializers.ValidationError({"organization_contribution": "La contribution dépasse le prix net."})
        return attrs


class ContractViewSet(ScopedModelViewSet):
    resource = "contracts"
    queryset = Contract.objects.select_related("organization", "training")
    serializer_class = ContractSerializer
    org_lookup = "organization"
    action_verbs = {"generate_invoices": "create"}
    filterset_fields = ["status", "organization", "training", "school"]
    search_fields = ["reference", "organization__name"]

    @action(detail=True, methods=["post"], url_path="generate-invoices")
    def generate_invoices(self, request, pk=None):
        """Facture l'organisation (sa contribution) et chaque bénéficiaire (part apprenant, en échéances)."""
        from apps.billing.services import generate_contract_invoices
        contract = self.get_object()
        first_due = request.data.get("first_due_date")
        first_due = date.fromisoformat(first_due) if first_due else date.today() + timedelta(days=30)
        result = generate_contract_invoices(contract, first_due)
        audit.log(request, "generate_invoices", contract, new=result)
        return Response(result, status=201)


class CohortSerializer(BaseModelSerializer):
    learners_count = serializers.SerializerMethodField()
    organization_name = serializers.CharField(source="organization.name", read_only=True, default=None)
    training_name = serializers.CharField(source="training.name", read_only=True)
    subscription = serializers.SerializerMethodField()

    class Meta:
        model = Cohort
        fields = "__all__"

    def get_learners_count(self, c):
        return c.learners.count()

    def get_subscription(self, c):
        """Plan de l'entreprise qui couvre cette cohorte (bénéficiaires => accès)."""
        from django.utils import timezone
        sub = (c.subscriptions.filter(status__in=["active", "suspended"], end_date__gte=timezone.localdate())
               .select_related("plan").order_by("-status", "end_date").first())
        if not sub:
            return None
        return {"plan_name": sub.plan.name, "status": sub.status, "end_date": sub.end_date, "seats": sub.seats,
                "reason": sub.suspension_reason}

    def validate(self, attrs):
        s, e = attrs.get("start_date"), attrs.get("end_date")
        if s and e and e < s:
            raise serializers.ValidationError({"end_date": "La fin doit suivre le début."})
        return attrs


def _seat_ok(cohort):
    from rest_framework.exceptions import ValidationError
    from apps.subscriptions.services import check_cohort_seat
    try:
        check_cohort_seat(cohort)
        return True
    except ValidationError:
        return False


HEADERS = {"nom": "last_name", "last_name": "last_name", "prenom": "first_name", "prénom": "first_name",
           "first_name": "first_name", "matricule": "employee_ref", "telephone": "phone", "téléphone": "phone",
           "phone": "phone", "email": "email", "e-mail": "email", "mail": "email"}


class CohortViewSet(ScopedModelViewSet):
    resource = "cohorts"
    queryset = Cohort.objects.select_related("organization", "training").prefetch_related("learners")
    serializer_class = CohortSerializer
    org_lookup = "organization"
    learner_lookup = "learners__user"
    instructor_lookup = "instructors__user"
    limit_key = "cohorts"
    parser_classes = [JSONParser, MultiPartParser, FormParser]
    action_verbs = {"import_learners": "create", "import_template": "view", "stats": "view"}
    filterset_fields = ["status", "organization", "training", "agency", "school"]
    search_fields = ["name", "location"]

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Nom", "Prénom", "Matricule", "Téléphone", "E-mail"])
        ws.append(["KOUASSI", "Awa", "EMP-001", "0700000000", "awa@exemple.ci"])
        resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = 'attachment; filename="modele_import_apprenants.xlsx"'
        wb.save(resp)
        return resp

    @action(detail=True, methods=["post"], url_path="import-learners", parser_classes=[MultiPartParser])
    def import_learners(self, request, pk=None):
        from openpyxl import load_workbook
        from apps.schools.services import check_limit
        cohort = self.get_object()
        f = request.FILES.get("file")
        if not f:
            return Response({"file": "Fichier requis."}, status=400)
        try:
            ws = load_workbook(f, read_only=True, data_only=True).active
            rows = list(ws.iter_rows(values_only=True))
        except Exception:
            return Response({"file": "Fichier Excel illisible."}, status=400)
        if not rows:
            return Response({"file": "Fichier vide."}, status=400)
        header = [HEADERS.get(str(h or "").strip().lower()) for h in rows[0]]
        if "last_name" not in header or "first_name" not in header:
            return Response({"file": "Colonnes 'Nom' et 'Prénom' obligatoires."}, status=400)
        created, errors = 0, []
        used = cohort.learners.count()
        emails = set(Learner.objects.filter(school=cohort.school).exclude(email="").values_list("email", flat=True))
        refs = set(Learner.objects.filter(school=cohort.school, organization=cohort.organization)
                   .exclude(employee_ref="").values_list("employee_ref", flat=True))
        with transaction.atomic():
            for idx, row in enumerate(rows[1:], start=2):
                rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row) if h}
                if not any(rec.values()):
                    continue
                if not rec.get("last_name") or not rec.get("first_name"):
                    errors.append({"row": idx, "message": "Nom ou prénom manquant."})
                elif rec.get("email") and rec["email"].lower() in emails:
                    errors.append({"row": idx, "message": "E-mail déjà utilisé."})
                elif rec.get("employee_ref") and rec["employee_ref"] in refs:
                    errors.append({"row": idx, "message": "Matricule déjà importé."})
                elif used >= cohort.capacity:
                    errors.append({"row": idx, "message": "Capacité de la cohorte atteinte."})
                elif not _seat_ok(cohort):
                    errors.append({"row": idx, "message": "Toutes les places de l'abonnement sont utilisées."})
                else:
                    check_limit(cohort.school, "learners")
                    Learner.objects.create(
                        school=cohort.school, organization=cohort.organization, cohort=cohort, training=cohort.training,
                        category=cohort.training.category, source="organization", status="registered", agency=cohort.agency,
                        email=rec.get("email", "").lower(), **{k: rec.get(k, "") for k in ("last_name", "first_name", "phone", "employee_ref")})
                    created += 1
                    used += 1
                    if rec.get("email"):
                        emails.add(rec["email"].lower())
                    if rec.get("employee_ref"):
                        refs.add(rec["employee_ref"])
        audit.log(request, "import_learners", cohort, new={"created": created, "errors": len(errors)})
        return Response({"created": created, "errors": errors, "total": created + len(errors)})

    @action(detail=True, methods=["get"])
    def stats(self, request, pk=None):
        from apps.reports.services import cohort_stats
        return Response(cohort_stats(self.get_object()))


class GroupViewSet(ScopedModelViewSet):
    resource = "groups"
    queryset = Group.objects.select_related("cohort")
    serializer_class = GroupSerializer
    school_lookup = "cohort__school"
    org_lookup = "cohort__organization"
    learner_lookup = "learners__user"
    instructor_lookup = "instructor__user"
    filterset_fields = ["cohort"]

    def perform_create(self, serializer):
        cohort = serializer.validated_data["cohort"]
        u = self.request.user
        if u.role != "akwaba_admin" and cohort.school_id != u.school_id:
            raise serializers.ValidationError({"cohort": "Cohorte inaccessible."})
        instance = serializer.save()
        audit.log(self.request, "create", instance, new=self._snapshot(instance))


QuoteRequestSerializer = make_serializer(QuoteRequest, read_only=("status",))


class QuoteRequestAdminSerializer(BaseModelSerializer):
    class Meta:
        model = QuoteRequest
        fields = "__all__"


class QuoteRequestViewSet(ScopedModelViewSet):
    """Demandes de devis de la vitrine : traitées par l'équipe Akwaba uniquement."""
    resource = "quotes"
    queryset = QuoteRequest.objects.all()
    serializer_class = QuoteRequestAdminSerializer
    school_lookup = None
    http_method_names = ["get", "patch", "delete", "head", "options"]
    filterset_fields = ["status"]
    search_fields = ["contact_name", "contact_email", "formation_label"]

    def get_queryset(self):
        from apps.core.permissions import AKWABA_ADMIN
        qs = QuoteRequest.objects.all()
        return qs if self.request.user.role == AKWABA_ADMIN else qs.none()


from rest_framework.permissions import AllowAny
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.views import APIView


class PublicQuoteView(APIView):
    """POST public (sans compte) depuis la vitrine."""
    authentication_classes = []
    permission_classes = [AllowAny]
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "auth"

    def post(self, request):
        ser = QuoteRequestSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        quote = ser.save()
        from apps.accounts.models import User
        from apps.core.notify import notify
        for admin in User.objects.filter(role="akwaba_admin", is_active=True):
            notify(admin, "quote_request", "Nouvelle demande de devis",
                   f"{quote.contact_name} - {quote.formation_label} ({quote.beneficiaries_count} bénéficiaires)")
        return Response({"detail": "Demande enregistrée."}, status=201)
