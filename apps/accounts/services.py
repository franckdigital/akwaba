"""Rattachement des comptes aux entreprises / cohortes et import Excel des utilisateurs."""
import os
import re
import secrets
import unicodedata

from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from rest_framework.exceptions import ValidationError

from apps.core.permissions import AKWABA_ADMIN, DIRECTOR, LEARNER, ORG_ADMIN, ORG_ROLES, ROLE_CHOICES
from apps.schools.models import CATEGORIES

from .models import User


# --------------------------------------------------------------------------- téléphone / connexion vocale
from .models import normalize_phone  # noqa: E402  (réexporté)


def users_by_phone(phone, voice_only=False):
    n = normalize_phone(phone)
    if len(n) < 8:
        return []
    qs = User.objects.filter(is_active=True, phone_norm=n)
    if voice_only:
        qs = qs.filter(voice_enabled=True)
    return list(qs)


VOICE_ROLES = {LEARNER, "instructor"}


def validate_voice_pin(pin, phone=""):
    """PIN vocal : 4 à 6 chiffres, ni répétition, ni suite, ni fin du numéro de téléphone."""
    pin = str(pin or "")
    if not re.fullmatch(r"\d{4,6}", pin):
        raise ValidationError({"pin": "Le code vocal doit contenir 4 à 6 chiffres."})
    if len(set(pin)) == 1 or pin in "0123456789" or pin in "9876543210":
        raise ValidationError({"pin": "Code trop simple (chiffres identiques ou suite) : choisissez-en un autre."})
    if phone and normalize_phone(phone).endswith(pin):
        raise ValidationError({"pin": "Le code ne doit pas être la fin de votre numéro de téléphone."})
    return pin


# --------------------------------------------------------------------------- rattachement
def link_membership(user, organization, cohort):
    """Relie un compte à une entreprise et, pour un apprenant, à une cohorte (dossier apprenant créé / mis à jour).

    Les bénéficiaires d'une cohorte couverte par un abonnement d'entreprise profitent automatiquement du plan.
    """
    from apps.learners.models import Learner
    from apps.schools.services import check_limit
    from apps.subscriptions.services import check_cohort_seat
    if cohort is not None and organization is not None and cohort.organization_id != organization.id:
        raise ValidationError({"cohort": "Cette cohorte n'appartient pas à l'entreprise choisie."})
    if cohort is not None and organization is None:
        organization = cohort.organization
    if organization is not None and organization.school_id != user.school_id:
        raise ValidationError({"organization": "Entreprise hors de l'auto-école du compte."})
    if cohort is not None and cohort.school_id != user.school_id:
        raise ValidationError({"cohort": "Cohorte hors de l'auto-école du compte."})
    user.organization = organization
    user.save(update_fields=["organization"])
    if user.role != LEARNER:
        return None
    learner = Learner.objects.filter(user=user).first()
    if learner is None:
        check_limit(user.school, "learners")
        learner = Learner(user=user, school=user.school, first_name=user.first_name, last_name=user.last_name or user.email,
                          email=user.email, phone=user.phone, category=cohort.training.category if cohort else "B", status="registered",
                          source="organization" if organization else "secretary")
    if cohort is not None and cohort != learner.cohort:
        check_cohort_seat(cohort, learner if learner.pk else None)
    learner.organization, learner.cohort = organization, cohort
    if cohort is None or (learner.group_id and learner.group.cohort_id != cohort.id):
        learner.group = None
    if cohort is not None and learner.training_id is None:
        learner.training = cohort.training
    learner.save()
    return learner


# --------------------------------------------------------------------------- import Excel
def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


HEADERS = {"prenom": "first_name", "first name": "first_name", "nom": "last_name", "last name": "last_name", "email": "email", "e mail": "email",
           "mail": "email", "telephone": "phone", "phone": "phone", "tel": "phone", "role": "role", "mot de passe": "password",
           "password": "password", "organisation": "organization", "entreprise": "organization", "etablissement": "organization",
           "organization": "organization", "cohorte": "cohort", "cohort": "cohort", "matricule": "employee_ref",
           "matricule employeur": "employee_ref", "permis": "category", "categorie": "category"}
ROLE_ALIASES = {_norm(label): key for key, label in ROLE_CHOICES}
ROLE_ALIASES.update({"apprenant": LEARNER, "beneficiaire": LEARNER, "etudiant": LEARNER, "employe": LEARNER, "administrateur d entreprise": ORG_ADMIN,
                     "admin entreprise": ORG_ADMIN, "rh": "hr", "secretaire": "secretary", "comptable": "accountant"})
FORBIDDEN_ROLES = {AKWABA_ADMIN, "instructor"}      # moniteurs : page dédiée (profil + catégories) ; admin Akwaba : jamais par import
EXAMPLES = [
    ["Awa", "KOUASSI", "awa.kouassi@exemple.ci", "0707123456", "Apprenant", "", "Entreprise ABC", "Cohorte Permis B - Septembre 2026", "EMP-001", "B"],
    ["Yves", "KONAN", "yves.konan@exemple.ci", "0505123456", "Apprenant", "Motdepasse#2026", "Entreprise ABC", "Cohorte Permis B - Septembre 2026", "EMP-002", "B"],
    ["Fatou", "DIALLO", "fatou.diallo@abc.ci", "0101123456", "RH", "", "Entreprise ABC", "", "", ""],
    ["Marc", "BAMBA", "marc.bamba@exemple.ci", "0707000000", "Comptable", "", "", "", "", ""],
]
INSTRUCTIONS = [
    ("Prénom", "Oui", "Prénom de l'utilisateur.", "Texte", "Awa"),
    ("Nom", "Oui", "Nom de famille.", "Texte", "KOUASSI"),
    ("E-mail", "Oui", "Adresse de connexion, unique.", "Adresse valide, jamais déjà utilisée", "awa@exemple.ci"),
    ("Téléphone", "Non", "Mobile (utilisé pour les paiements Mobile Money).", "10 chiffres", "0707123456"),
    ("Rôle", "Non (défaut : Apprenant)", "Rôle de l'utilisateur.", "Apprenant, Secrétaire, Comptable, Directeur, RH, Responsable formation, Administrateur d'entreprise, Comptable entreprise, Direction", "Apprenant"),
    ("Mot de passe", "Non", "Laissé vide : un mot de passe temporaire est généré et affiché une seule fois après l'import.", "8 caractères minimum", ""),
    ("Organisation", "Oui pour les rôles d'entreprise", "Nom EXACT d'une entreprise / d'un établissement existant.", "Voir l'onglet « Organisations & cohortes »", "Entreprise ABC"),
    ("Cohorte", "Non (apprenants)", "Nom EXACT d'une cohorte de l'organisation : l'apprenant en bénéficie et profite du plan souscrit pour cette cohorte.", "Voir l'onglet « Organisations & cohortes »", "Cohorte Permis B - Septembre 2026"),
    ("Matricule employeur", "Non", "Matricule chez l'employeur / l'établissement.", "Texte", "EMP-001"),
    ("Permis", "Non (défaut : celui de la cohorte, sinon B)", "Catégorie visée.", "A, B, C, D, E", "B"),
]


def _initial_learner_room(school):
    """Places restantes dans le forfait (None = illimité)."""
    from apps.schools.models import Subscription
    from apps.schools.services import active_subscription, count_for
    if not Subscription.objects.filter(school=school).exists():
        return None
    sub = active_subscription(school)
    if sub is None:
        return 0
    return max(sub.plan.max_learners - count_for(school, "learners"), 0)


def build_template(school, blank=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from apps.organizations.models import Cohort, Organization
    wb = Workbook()
    ws = wb.active
    ws.title = "Utilisateurs"
    ws.append(["Prénom", "Nom", "E-mail", "Téléphone", "Rôle", "Mot de passe", "Organisation", "Cohorte", "Matricule employeur", "Permis"])
    red, grey = PatternFill("solid", fgColor="E02027"), PatternFill("solid", fgColor="2B2B2B")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = red if c.value in ("Prénom", "Nom", "E-mail") else grey
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if not blank:
        for r in EXAMPLES:
            ws.append(r)
    for col, w in zip("ABCDEFGHIJ", (14, 16, 30, 14, 16, 18, 26, 36, 18, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("Mode d'emploi")
    ins.append(["IMPORT DES UTILISATEURS — MODE D'EMPLOI"])
    ins["A1"].font = Font(bold=True, size=14, color="E02027")
    ins.append([])
    ins.append(["Colonne", "Obligatoire", "Description", "Valeurs acceptées", "Exemple"])
    for c in ins[3]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), red
    for r in INSTRUCTIONS:
        ins.append(list(r))
    ins.append([])
    ins.append(["Règles : une ligne = un utilisateur ; les lignes d'exemple sont à supprimer avant import ; le fichier est ANALYSÉ avant tout enregistrement ; les mots de passe générés ne sont affichés qu'une fois."])
    for col, w in zip("ABCDE", (20, 26, 70, 50, 34)):
        ins.column_dimensions[col].width = w
    for row in ins.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ref = wb.create_sheet("Organisations & cohortes")
    ref.append(["Organisation", "Cohorte"])
    for c in ref[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), grey
    if school is not None:
        for org in Organization.objects.filter(school=school):
            ref.append([org.name, ""])
            for co in Cohort.objects.filter(organization=org):
                ref.append([org.name, co.name])
    ref.column_dimensions["A"].width, ref.column_dimensions["B"].width = 34, 44
    return wb


def _allowed_roles(actor):
    if actor.role == AKWABA_ADMIN:
        return {k for k, _ in ROLE_CHOICES} - FORBIDDEN_ROLES
    if actor.role == DIRECTOR:
        return {k for k, _ in ROLE_CHOICES} - FORBIDDEN_ROLES - {AKWABA_ADMIN}
    if actor.role == ORG_ADMIN:
        return set(ORG_ROLES) | {LEARNER}
    return set()


def run_import(fileobj, actor, school, commit):
    """Analyse (commit=False) ou crée les comptes d'un fichier Excel. -> rapport (avec identifiants générés si commit)."""
    from openpyxl import load_workbook
    from apps.learners.models import Learner
    from apps.organizations.models import Cohort, Organization
    from apps.subscriptions.models import OrganizationSubscription
    from apps.subscriptions.services import today
    if school is None:
        raise ValidationError({"school": "Auto-école requise."})
    try:
        rows = list(load_workbook(fileobj, read_only=True, data_only=True).active.iter_rows(values_only=True))
    except Exception:
        raise ValidationError({"file": "Fichier Excel (.xlsx) illisible."})
    if not rows:
        raise ValidationError({"file": "Fichier vide."})
    if len(rows) > 1001:
        raise ValidationError({"file": "1 000 lignes maximum par import : scindez le fichier."})
    header = [HEADERS.get(_norm(h)) for h in rows[0]]
    missing = {"first_name", "last_name", "email"} - set(header)
    if missing:
        raise ValidationError({"file": "Colonnes manquantes : " + ", ".join(sorted(missing)) + ". Utilisez le modèle Excel."})
    allowed = _allowed_roles(actor)
    orgs = {_norm(o.name): o for o in Organization.objects.filter(school=school)}
    existing = set(User.objects.values_list("email", flat=True))
    seen, errors, valid, credentials, created = set(), [], [], [], 0
    to_create = []
    room = {}    # places restantes par cohorte couverte (simulation pour l'analyse à blanc)
    total = 0

    limit_room = [None]
    limit_room[0] = _initial_learner_room(school)

    def learner_room():
        return limit_room[0]

    def seat_room(cohort):
        if cohort.id not in room:
            subs = OrganizationSubscription.objects.filter(cohorts=cohort, status="active", start_date__lte=today(), end_date__gte=today())
            vals = [s.seats - Learner.objects.filter(cohort__in=s.cohorts.all()).count() for s in subs if s.seats]
            room[cohort.id] = min(vals) if vals else None
        return room[cohort.id]

    with transaction.atomic():
        for idx, row in enumerate(rows[1:], start=2):
            if row is None or all(v in (None, "") for v in row):
                continue
            total += 1
            rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row) if h}
            email = rec.get("email", "").lower()
            problems = []
            if not rec.get("first_name") or not rec.get("last_name"):
                problems.append("Prénom et nom obligatoires.")
            if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
                problems.append("E-mail invalide.")
            elif email in existing or email in seen:
                problems.append("E-mail déjà utilisé.")
            role = ROLE_ALIASES.get(_norm(rec.get("role"))) if rec.get("role") else LEARNER
            if role is None:
                problems.append(f"Rôle inconnu « {rec.get('role')} ».")
            elif role not in allowed:
                problems.append(f"Rôle « {rec.get('role') or role} » non autorisé pour cet import.")
            org = cohort = None
            if rec.get("organization"):
                org = orgs.get(_norm(rec["organization"]))
                if org is None:
                    problems.append(f"Organisation « {rec['organization']} » introuvable.")
            if actor.role == ORG_ADMIN:
                org = actor.organization
            if role in ORG_ROLES and org is None and not problems:
                problems.append("Organisation obligatoire pour un rôle d'entreprise.")
            if rec.get("cohort") and role == LEARNER:
                qs = Cohort.objects.filter(school=school, name__iexact=rec["cohort"])
                if org is not None:
                    qs = qs.filter(organization=org)
                found = list(qs[:2])
                if not found:
                    problems.append(f"Cohorte « {rec['cohort']} » introuvable" + (f" pour {org.name}." if org else "."))
                elif len(found) > 1:
                    problems.append(f"Cohorte « {rec['cohort']} » ambiguë : précisez l'organisation.")
                else:
                    cohort = found[0]
                    org = org or cohort.organization
            elif rec.get("cohort"):
                problems.append("Une cohorte ne concerne que les apprenants.")
            category = (rec.get("category") or "").upper()
            if category and category not in {c for c, _ in CATEGORIES}:
                problems.append(f"Permis invalide « {rec.get('category')} ».")
            pwd = rec.get("password") or ""
            if pwd:
                try:
                    validate_password(pwd)
                except DjangoValidationError as e:
                    problems.append("Mot de passe : " + " ".join(e.messages))
            if not problems and role == LEARNER:
                if cohort is not None:
                    r = seat_room(cohort)
                    if r is not None and r <= 0:
                        problems.append("Toutes les places de l'abonnement de cette cohorte sont utilisées.")
                    elif r is not None:
                        room[cohort.id] = r - 1
                if not problems:
                    r = learner_room()
                    if r is not None and r <= 0:
                        problems.append("Limite d'apprenants du forfait atteinte.")
                    elif r is not None:
                        limit_room[0] = r - 1
            if problems:
                errors.append({"row": idx, "email": email, "message": " ".join(problems), "data": [("" if v is None else str(v)) for v in row]})
                continue
            seen.add(email)
            item = {"row": idx, "name": f"{rec['first_name']} {rec['last_name']}", "email": email, "role": role,
                    "organization": org.name if org else "", "cohort": cohort.name if cohort else ""}
            valid.append(item)
            if commit:
                to_create.append((item, rec, org, cohort, role, pwd or secrets.token_urlsafe(9), bool(pwd), category))
        if commit and to_create:
            # PBKDF2 est volontairement lent (~1 s / mot de passe) : hachage en parallèle (hashlib libère le GIL)
            from concurrent.futures import ThreadPoolExecutor
            from django.contrib.auth.hashers import make_password
            with ThreadPoolExecutor(max_workers=min(8, os.cpu_count() or 2)) as pool:
                hashes = list(pool.map(make_password, [t[5] for t in to_create]))
            for (item, rec, org, cohort, role, password, given, category), hashed in zip(to_create, hashes):
                user = User(email=item["email"], first_name=rec["first_name"], last_name=rec["last_name"], phone=rec.get("phone", ""),
                            role=role, school=school, password=hashed)
                user.save()
                learner = link_membership(user, org, cohort)
                if learner is not None:
                    if rec.get("employee_ref"):
                        learner.employee_ref = rec["employee_ref"]
                    if category:
                        learner.category = category
                    learner.save()
                credentials.append({"email": item["email"], "name": item["name"], "role": role, "password": password, "generated": not given})
                created += 1
        if not commit:
            transaction.set_rollback(True)
    return {"total": total, "valid": len(valid), "created": created, "error_count": len(errors), "errors": errors[:300],
            "preview": valid[:30], "credentials": credentials, "dry_run": not commit}
